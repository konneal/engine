#!/usr/bin/env python3
"""Convert the Bulletin article (MDX) to a single-column .docx submission.

Usage: python3 scripts/mdx_to_docx.py docs/paper-oiml-bulletin.mdx -o artifacts/paper/

Handles the MDX subset the article uses: frontmatter, headings, paragraphs
with **bold** / *italic* / `code` inline, bullet and numbered lists, pipe
tables, blockquotes, horizontal rules, and <Figure src="..." caption="..."/>
JSX blocks (SVG figures are rasterized with rsvg-convert and embedded; PNG
twins are written beside the .docx for separate-file submission).
"""
import argparse
import re
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

FIGURE_RE = re.compile(r"<Figure\s*([^/]*?)\s*/>", re.DOTALL)
ATTR_RE = re.compile(r'(\w+)\s*=\s*"([^"]*)"')


def parse_figure(block: str) -> dict:
    attrs = dict(ATTR_RE.findall(block))
    return attrs  # src, caption, id


def rasterize(svg_path: Path, png_path: Path, width_px: int = 1920) -> bool:
    if shutil.which("rsvg-convert") is None:
        return False
    png_path.parent.mkdir(parents=True, exist_ok=True)
    subprocess.run(
        ["rsvg-convert", "-w", str(width_px), "-o", str(png_path), str(svg_path)],
        check=True,
    )
    return True


INLINE_TOKEN = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`|\[[^\]]+\]\([^)]+\))")


def add_inline(paragraph, text: str):
    for part in INLINE_TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Courier New"
        elif part.startswith("[") and "](" in part:
            label, url = part[1:-1].split("](", 1)
            paragraph.add_run(label)
            run = paragraph.add_run(f" ({url})")
            run.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
        else:
            paragraph.add_run(part)


def strip_frontmatter(text: str):
    if text.startswith("---\n"):
        end = text.find("\n---\n", 4)
        if end != -1:
            meta = text[4 : end]
            title = ""
            for line in meta.splitlines():
                if line.startswith("title:"):
                    title = line.split("title:")[1].strip().strip('"')
            return text[end + 5 :], title
    return text, ""


def strip_mdx_exports(text: str) -> str:
    # drop `export const Figure = ( ... );` blocks — layout code, not content
    return re.sub(r"^export const .*?;\n", "", text, flags=re.DOTALL | re.MULTILINE)


def convert(mdx_path: Path, out_dir: Path) -> Path:
    text = mdx_path.read_text(encoding="utf-8")
    text, title = strip_frontmatter(text)
    text = strip_mdx_exports(text)

    doc = Document()
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Cm(2.5)
        section.left_margin = section.right_margin = Cm(2.5)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)
    if title:
        doc.core_properties.title = title

    out_dir.mkdir(parents=True, exist_ok=True)
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]

        fig = FIGURE_RE.match(line.strip()) or (
            FIGURE_RE.match("\n".join(lines[i : i + 6]).strip()) if line.strip().startswith("<Figure") else None
        )
        if line.strip().startswith("<Figure"):
            block = [line]
            while not block[-1].strip().endswith("/>") and i + len(block) < len(lines):
                block.append(lines[i + len(block)])
            attrs = parse_figure("\n".join(block))
            i += len(block)
            src = mdx_path.parent / attrs.get("src", "")
            caption = attrs.get("caption", "")
            if src.suffix == ".svg":
                png = out_dir / (src.stem + ".png")
                if rasterize(src, png):
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.add_run().add_picture(str(png), width=Cm(15.5))
                    cap = doc.add_paragraph()
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    add_inline(cap, caption)
                    for run in cap.runs:
                        run.font.size = Pt(9)
                    continue
            doc.add_paragraph(caption)  # figure fell back to caption-only
            continue

        if line.startswith("|") and i + 1 < len(lines) and set(lines[i + 1].replace("|", "").replace("-", "").strip()) <= {":", " "}:
            rows = [r for r in lines[i : i + 20] if r.startswith("|")]
            header = [c.strip() for c in rows[0].strip("|").split("|")]
            body = [[c.strip() for c in r.strip("|").split("|")] for r in rows[2:]]
            i += len(rows[: 2 + len(body)])
            table = doc.add_table(rows=1 + len(body), cols=len(header))
            table.style = "Table Grid"
            for j, cell in enumerate(header):
                para = table.rows[0].cells[j].paragraphs[0]
                add_inline(para, cell)
                for run in para.runs:
                    run.bold = True
            for r, row in enumerate(body, start=1):
                for j, cell in enumerate(row[: len(header)]):
                    add_inline(table.rows[r].cells[j].paragraphs[0], cell)
            continue

        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped == "---":
            i += 1
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            doc.add_heading(m.group(2), level=min(len(m.group(1)), 4))
            i += 1
            continue
        if stripped.startswith("> "):
            p = doc.add_paragraph(style="Intense Quote")
            add_inline(p, stripped[2:])
            i += 1
            continue
        def is_block_start(s: str) -> bool:
            return (
                not s
                or s.startswith(("#", "|", ">", "<Figure", "---"))
                or re.match(r"^([-*]|\d+\.)\s", s) is not None
            )

        def absorb_continuation() -> tuple[str, int]:
            extra = []
            while i + 1 + len(extra) < len(lines) and not is_block_start(lines[i + 1 + len(extra)].strip()):
                extra.append(lines[i + 1 + len(extra)].strip())
            return " ".join(extra), len(extra)

        if re.match(r"^[-*]\s+", stripped):
            extra, n = absorb_continuation()
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, (re.sub(r"^[-*]\s+", "", stripped) + " " + extra).strip())
            i += 1 + n
            continue
        if re.match(r"^\d+\.\s+", stripped):
            extra, n = absorb_continuation()
            p = doc.add_paragraph(style="List Number")
            add_inline(p, (re.sub(r"^\d+\.\s+", "", stripped) + " " + extra).strip())
            i += 1 + n
            continue

        para_lines = [line]
        while i + len(para_lines) < len(lines):
            nxt = lines[i + len(para_lines)].strip()
            if not nxt or nxt.startswith(("#", "|", "-", ">", "<Figure", "---")) or re.match(r"^\d+\.\s", nxt):
                break
            para_lines.append(lines[i + len(para_lines)])
        p = doc.add_paragraph()
        add_inline(p, " ".join(l.strip() for l in para_lines))
        i += len(para_lines)

    out = out_dir / (mdx_path.stem + ".docx")
    doc.save(out)
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("mdx", type=Path)
    ap.add_argument("-o", "--out", type=Path, default=Path("artifacts/paper"))
    args = ap.parse_args()
    out = convert(args.mdx, args.out)
    print(f"wrote {out}")
    for png in sorted(args.out.glob("*.png")):
        print(f"figure file: {png}")


if __name__ == "__main__":
    sys.exit(main())
